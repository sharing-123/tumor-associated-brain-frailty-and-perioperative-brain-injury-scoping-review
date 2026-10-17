import sys, json, re
sys.stdout.reconfigure(encoding="utf-8")
BUILD = "."
sys.path.insert(0, BUILD)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from data_core import *
from prisma_numbers import P
import stats
from tables import FLAGNAME
import content_b as B

OUT = "../outputs/Supplementary_Tables_S1-S6.xlsx"
REFS = json.load(open(BUILD + r"\refs_master.json", encoding="utf-8"))
ZK = json.load(open(BUILD + "/zkeys.json"))
ORDER = json.load(open(BUILD + r"\cite_order.json"))
CNUM = {k: i + 1 for i, k in enumerate(ORDER)}
CR = json.load(open(BUILD + r"\cr_all.json", encoding="utf-8"))
XL = openpyxl.load_workbook("../inputs/Esktraksi Data (1).xlsx")["Master Extraction"]
XROWS = [r for r in XL.iter_rows(min_row=2, values_only=True) if r[0]]
XH = [c.value for c in XL[1]]

HF = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3B57")
thin = Side(style="thin", color="BBBBBB")

def sheet(wb, name, header, rows, widths=None, wrap=True):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = HF
        c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=wrap, vertical="top")
    for i, h in enumerate(header):
        w = (widths[i] if widths else min(max(len(str(h)) + 2, 12), 40))
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"
    return ws

wb = openpyxl.Workbook()
wb.remove(wb.active)

s1 = [
 (1, "Title", "Identify the report as a scoping review", "Title page", "Reported"),
 (2, "Structured summary", "Structured summary with background, objectives, eligibility, sources, charting, results, conclusions", "Abstract", "Reported"),
 (3, "Rationale", "Describe the rationale in the context of what is already known", "1. Introduction", "Reported"),
 (4, "Objectives", "Provide an explicit statement of questions and objectives using PCC", "1. Introduction (RQ1 to RQ7), 2.3", "Reported"),
 (5, "Protocol and registration", "Indicate whether a protocol exists, where it can be accessed, and registration details", "2.1", "Reported (not registered, materials deposited on GitHub, Zenodo archive to follow)"),
 (6, "Eligibility criteria", "Specify characteristics of the sources of evidence used as eligibility criteria", "2.3", "Reported"),
 (7, "Information sources", "Describe all information sources with the date last searched", "2.4 (last searches 22 March 2026)", "Reported"),
 (8, "Search", "Present the full electronic search strategy for at least one database", "2.4 and Supplementary Table S2", "Reported"),
 (9, "Selection of sources of evidence", "State the process for selecting sources of evidence", "2.5 (two reviewers screened independently, disagreements resolved in a video meeting)", "Reported"),
 (10, "Data charting process", "Describe the methods of charting data and any processes for confirming data", "2.6 and Supplementary Table S3", "Reported"),
 (11, "Data items", "List and define all variables for which data were sought", "2.6", "Reported"),
 (12, "Critical appraisal of individual sources", "If done, provide the rationale and methods used", "2.8 (not performed, design-based evidence level used)", "Reported as not done"),
 (13, "Synthesis of results", "Describe the methods of handling and summarizing the data", "2.7", "Reported"),
 (14, "Selection of sources of evidence", "Give numbers screened, assessed, and included with reasons and a flow diagram", "3.1 and Figure 1", "Reported"),
 (15, "Characteristics of sources of evidence", "Present characteristics for which data were charted", "3.2, Table 1, Figure 2", "Reported"),
 (16, "Critical appraisal within sources", "If done, present appraisal data", "Not applicable", "Not done"),
 (17, "Results of individual sources of evidence", "Present the relevant data charted for each source", "Table 1 and Supplementary Table S3", "Reported"),
 (18, "Synthesis of results", "Summarize and link results to the review questions", "3.3 to 3.9, Tables 2 to 7, Figures 3 to 5", "Reported"),
 (19, "Summary of evidence", "Summarize main results and relevance to the target groups", "4. Discussion", "Reported"),
 (20, "Limitations", "Discuss the limitations of the scoping review process", "4. Discussion", "Reported"),
 (21, "Conclusions", "Provide a general interpretation and potential implications", "5. Conclusions", "Reported"),
 (22, "Funding", "Describe sources of funding and the role of funders", "Back matter (no external funding)", "Reported"),
]
sheet(wb, "S1 PRISMA-ScR", ["Item", "Section", "Checklist item", "Where reported", "Status"], s1, [6, 28, 70, 50, 26])

d = docx.Document("../inputs/V2_Hasil Searching.docx")
rows = []
for r in d.tables[0].rows[1:]:
    c = [x.text.strip() for x in r.cells]
    rows.append([c[0], c[1], c[2]])
rows.append(["Citation searching", "Reference lists and citing articles of included studies and relevant reviews", str(P["citation"])])
rows.append(["Client-supplied sources", "Sources on MoCA-INA and NANO-Ina supplied by the clinical team (Zoom consultation and links)", str(P["client"])])
sheet(wb, "S2 Searches", ["Source", "Search strategy", "Records (n)"], rows, [26, 140, 14])

hdr = XH[:30] + ["Reference key (Zotero)", "Citation number in manuscript", "Evidence stream", "Design class", "OCEBM level", "Tumor category", "Domain codes", "Verified DOI", "PMID", "Publication year used in manuscript"]
rows = []
for r in XROWS:
    no = r[0]
    k = "core%d" % no
    ref = REFS[k]
    base = list(r[:30])
    rows.append(base + [ZK["keys"][k], CNUM[k], STREAMS[R[no][4]], R[no][2], level(no), R[no][3], "".join(R[no][9]), ref["doi"] or "No DOI exists (PMID %s)" % ref.get("pmid", ""), ref.get("pmid", ""), stats.YEAR[no]])
for k, (lab, ctry, des, pop, aim, inst, fnd) in NEW.items():
    ref = REFS[k]
    blank = [None] * 30
    blank[0] = "I-%d" % (list(NEW).index(k) + 1)
    blank[1] = lab.split(",")[0]
    blank[2] = lab.split(", ")[1]
    blank[3] = ref["journal"]
    blank[4] = ctry
    blank[5] = des
    blank[6] = pop
    blank[8] = pop
    blank[19] = inst
    blank[28] = fnd
    blank[29] = ref["doi"]
    rows.append(blank + [ZK["keys"][k], CNUM[k], "Instrument-focused stratum", des, "Not graded", "Instrument", "", ref["doi"], ref.get("pmid", ""), ref["year"]])
sheet(wb, "S3 Charting form", hdr, rows, [8, 18, 8, 22, 14, 24, 28, 28, 22, 18, 14, 28, 18, 28, 20, 24, 20, 18, 18, 22, 22, 22, 22, 22, 22, 22, 22, 22, 70, 40, 12, 14, 14, 24, 10, 16, 10, 18, 10, 10, 10])

ws = wb.create_sheet("S4 Reconciliation")
ws.append(["A. Screening numbers, original versus final"])
ws.append(["Step", "Original tally (80 included)", "Final (76 included)", "Note"])
orig = [
 ("Databases (PubMed, Scopus, ProQuest, Cochrane)", 2302, P["db_only"], "The original template labelled 2,318 as Databases. The four databases sum to 2,302, and the other 16 are the medRxiv preprint server."),
 ("Preprint server (medRxiv)", 16, P["medrxiv"], "Listed as Registers in the template. medRxiv is a preprint server, not a trial register."),
 ("Google Scholar", 117, P["gscholar"], ""),
 ("Citation searching", 34, P["citation"], ""),
 ("Client-supplied instrument literature", 0, P["client"], "Added after consultation. Seven records."),
 ("Duplicates removed (database and preprint arm)", 987, P["dups"], ""),
 ("Records screened", 1331, P["screened"], "2,318 minus 987"),
 ("Excluded at title and abstract", 1062, P["ta_excl"], ""),
 ("Reports sought, database arm", 269, P["sought_db"], ""),
 ("Not retrieved, database arm", 17, P["nr_db"], ""),
 ("Assessed, database arm", 252, P["assessed_db"], ""),
 ("Excluded, database arm", 206, P["excl_db"], "Nine studies added from the charting stage"),
 ("Included, database arm", 46, P["inc_db"], ""),
 ("Reports sought, other methods", 151, P["sought_ot"], "117 plus 34 originally, 158 with the 7 client records"),
 ("Not retrieved, other methods", 13, P["nr_ot"], ""),
 ("Assessed, other methods", 138, P["assessed_ot"], ""),
 ("Excluded, other methods", 104, P["excl_ot"], "One study from charting and one instrument manual added"),
 ("Included, other methods", 34, P["inc_ot"], "33 core plus 6 instrument-focused studies"),
 ("Total included", 80, P["inc_total"], "70 core (extraction workbook) plus 6 instrument-focused"),
]
for r in orig:
    ws.append(list(r))
ws.append([])
ws.append(["B. Ten studies removed during data charting (in the 80, not in the 70-row extraction workbook)"])
ws.append(["Arm", "First author, year", "Title", "DOI", "Reason assigned", "Confirm with team"])
removed = [
 ("Other methods", "Muller, 2017", "Single-cell profiling of human gliomas reveals macrophage ontogeny as a basis for regional differences in macrophage activation in the tumor microenvironment", "10.1186/s13059-017-1362-4", "Wrong outcome (molecular origin of macrophages, no perioperative or clinical outcome)"),
 ("Database", "Hu, 2010", "Effects of remote ischemic preconditioning on biochemical markers and neurologic outcomes in patients undergoing elective cervical decompression surgery (file named Kutum)", "J Neurosurg Anesthesiol 2010;22:46-52", "Wrong population (cervical spondylotic myelopathy, spine surgery)"),
 ("Database", "Zhang, 2022", "Analysis of cerebrospinal fluid routine biochemical level, pathogenic bacteria distribution, and risk factors in patients with secondary intracranial infection after brain tumor surgery", "10.1155/2022/7716205", "Wrong outcome (CSF infection biomarkers and pathogen distribution, no frailty, brain injury, or cognitive outcome). Borderline, excluded"),
 ("Database", "Del Gaudio, 2006", "Remifentanil vs fentanyl with a target controlled propofol infusion in patients undergoing craniotomy for supratentorial lesions", "Minerva Anestesiol 2006;72:309-19", "Wrong outcome (hemodynamics and recovery time only)"),
 ("Database", "Sanai, 2018", "Phase 0 trial of AZD1775 in first-recurrence glioblastoma patients", "10.1158/1078-0432.CCR-17-3348", "Wrong outcome (pharmacokinetic phase 0 drug trial, no perioperative brain injury or cognitive outcome)"),
 ("Database", "Sonabend, 2023", "Repeated blood-brain barrier opening by an implantable ultrasound device for delivery of albumin-bound paclitaxel, a phase 1 trial in recurrent glioblastoma", "10.1016/S1470-2045(23)00112-2", "Wrong outcome (SonoCloud-9 drug delivery trial, no perioperative or cognitive outcome)"),
 ("Database", "Kan, 2025", "Characterisation of the expression of P2X7 receptor, cancer stem cell markers and immunological mediators in human high-grade gliomas", "10.1186/s12868-025-00973-5", "Wrong outcome (molecular characterization of P2X7 receptor in glioma tissue, no clinical outcome)"),
 ("Database", "Carpentier, 2024", "Repeated blood-brain barrier opening with a nine-emitter implantable ultrasound device in combination with carboplatin in recurrent glioblastoma", "10.1038/s41467-024-45818-7", "Wrong outcome (SonoCloud-9 drug delivery trial, no perioperative or cognitive outcome)"),
 ("Database", "Magrassi, 2023", "Proteomic analysis on sequential samples of cystic fluid obtained from human brain tumors", "10.3390/cancers15164070", "Wrong outcome (cyst fluid proteomics)"),
 ("Database", "Doniselli, 2025", "Cerebral microbubble distribution correlation between preoperative perfusion MRI and intraoperative contrast enhanced ultrasound", "10.1038/s44384-025-00035-4", "Wrong outcome (comparison of CEUS and DSC MRI, no patient clinical outcome)"),
]
for r in removed:
    ws.append(list(r) + ["Confirmed by the review team"])
ws.append(["Note", "Yuan 2023", "The database folder held two PDF copies of Yuan 2023 (the FORTITUDE trial). The duplicate file was dropped and the study stays in the core set as study 67. The count of 10 removed studies therefore includes Del Gaudio 2006, which is absent from the extraction workbook. Its reason (wrong outcome) was confirmed by the team."])
ws.append([])
ws.append(["C. Client-supplied records and decisions"])
ws.append(["Record", "Type", "Decision", "Reason", "Used for"])
cs = [
 ("Nayak 2017, NANO scale (PubMed 28453751)", "Primary instrument study", "Included, instrument-focused stratum", "Neuro-oncology population, neurologic function instrument", "Neurological domain"),
 ("Angelita 2026, NANO-Ina (PubMed 41793672)", "Validation study", "Included, instrument-focused stratum", "Indonesian version, brain tumor patients", "Neurological domain"),
 ("Jammula 2022, MoCA in neuro-oncology", "Cross-sectional feasibility study", "Included, instrument-focused stratum", "Primary CNS tumor patients, parent instrument", "Cognitive domain"),
 ("Bangun 2018, MMSE versus MoCA-Ina in brain tumors", "Cross-sectional study (book chapter)", "Included, instrument-focused stratum", "Indonesian brain tumor patients at diagnosis", "Cognitive domain"),
 ("Rambe 2017, MoCA-INA versus MMSE in elderly", "Cross-sectional study", "Included, instrument-focused stratum", "No tumor patients, Indonesian version in adults (language and population criteria relaxed)", "Cognitive domain, instrument context"),
 ("Wonggo 2025, MoCA-INA in anesthesiology residents", "Cross-sectional study", "Included, instrument-focused stratum", "No tumor patients, Indonesian version in adults, Indonesian text with English abstract", "Cognitive domain, instrument context"),
 ("MoCA-InA instruction and scoring sheet", "Instrument manual, no data", "Excluded as a study (wrong publication type)", "No study data", "Scoring rules in Results 3.8 (maximum 30, 26 or more normal, plus 1 point for 12 or fewer years of education)"),
]
for r in cs:
    ws.append(list(r))
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for i, w in enumerate([46, 24, 60, 42, 52, 30]):
    ws.column_dimensions[get_column_letter(i + 1)].width = w
for rr in (1, 2):
    pass
for r in ws.iter_rows():
    if r[0].value and str(r[0].value).startswith(("A.", "B.", "C.")):
        r[0].font = Font(bold=True, size=12)
    if r[0].value in ("Step", "Arm", "Record"):
        for c in r:
            c.font = HF
            c.fill = HFILL

hdr5 = ["Citation no.", "Reference key", "Zotero item key", "Type", "First author", "Year", "Journal or book", "Journal abbreviation (NLM)", "Volume", "Issue", "Pages", "DOI", "PMID", "Crossref match", "Note"]
notes = {
 "core8": "DOI in extraction workbook (10.1097/ANA.0b013e3181c9ea63) did not resolve. Corrected to 10.1097/ANA.0b013e3181c9fbf1 after Crossref title and author match.",
 "core50": "DOI in extraction workbook (10.1007/s11655-010-029-6) did not resolve. Corrected to 10.1007/s11655-010-0229-6.",
 "core42": "DOI in extraction workbook was truncated. Full DOI resolved through Crossref title and author match.",
 "core55": "Published 2016 (issue 28:1), the workbook listed 2015.",
 "core12": "No DOI exists. Identified by PubMed (PMID 15140713).",
 "core13": "No DOI exists. Identified by PubMed (PMID 16484420).",
 "core60": "Preprint, not peer reviewed. DOI resolved through DataCite.",
 "core39": "Page range confirmed in PubMed (45-57).",
 "appleton2020": "Page range confirmed in PubMed (e439-e452).",
 "core32": "Title garbled in Crossref, cleaned from the article PDF.",
 "bangun2018": "Book chapter. DOI resolved through Crossref.",
 "wonggo2025": "Indonesian-language article with English title and abstract. DOI resolved through Crossref.",
}
rows = []
for k in ORDER:
    r = REFS[k]
    rows.append([CNUM[k], k, ZK["keys"][k], r["type"], r["authors"][0]["family"] if r["authors"] else "", r["year"], r["journal"], r["abbr"], r["volume"], r["issue"], r["pages"], r["doi"] or "None (see note)", r.get("pmid", ""), "Yes (DOI resolved in Crossref and fields compared)" if r["doi"] and k not in ("core60",) else ("DataCite" if k == "core60" else "PubMed"), notes.get(k, "")])
sheet(wb, "S5 Reference log", hdr5, rows, [8, 14, 12, 14, 16, 7, 34, 24, 8, 8, 14, 34, 11, 30, 70])

ws5 = wb["S5 Reference log"]
ws5.append([])
ws5.append(["Discrepancies between the extraction workbook and verified records"])
ws5.append(["Study no.", "Workbook label", "Workbook year", "Verified year", "Workbook DOI or locator", "Verified DOI"])
for r in XROWS:
    no = r[0]
    ref = REFS["core%d" % no]
    wy = int(str(r[2])[:4])
    m = re.search(r"10\.\d{4,9}/\S+", str(r[29]))
    wdoi = m.group(0).rstrip(".)") if m else str(r[29])
    if wy != stats.YEAR[no] or (ref["doi"] and wdoi.lower() != ref["doi"].lower()):
        ws5.append([no, r[1], wy, stats.YEAR[no], wdoi, ref["doi"] or "No DOI exists"])

rows = []
for no in sorted(R):
    lab, ctry, des, tum, st, n, pop, ctx, fnd, fl = R[no]
    rows.append([no, lab, ctry, des, level(no), tum, STREAMS[st], n, ", ".join(FLAGNAME[f] for f in fl), "Yes" if no in B.F_IDS else "", "Yes" if no in B.C_IDS else "", "Yes" if no in B.N_IDS else "", "Yes" if no in B.I_IDS else ""])
sheet(wb, "S6 Coding", ["No.", "Study", "Country", "Design class", "OCEBM level (design only)", "Tumor category", "Evidence stream", "Sample size", "Domains", "Domain 1 systemic frailty", "Domain 2 cognition", "Domain 3 neurological", "Domain 4 imaging"], rows, [6, 26, 14, 26, 12, 14, 44, 10, 50, 12, 12, 12, 12])
wb.save(OUT)
print("saved", OUT)
