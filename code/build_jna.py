import os, re, sys, json, zipfile, shutil
sys.path.insert(0, ".")
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import openpyxl
import jna_lib as L
import content_jna as J
import tables as T

OUT = "../outputs/JNA_submission"
FIGS = "figs"
os.makedirs(OUT, exist_ok=True)
FONT = L.FONT
CSL_ID = "https://github.com/sharing-123/tumor-associated-brain-frailty-and-perioperative-brain-injury-scoping-review/jna-style"

MAIN_TEXTS = list(J.INTRO) + [p for k, p in J.METHODS if k == "p"] + [p for k, p in J.RESULTS if k == "p"] + list(J.DISCUSSION) + [J.CONCLUSIONS]
MAIN_ORDER = L.order_of(MAIN_TEXTS)
MAIN = L.Ctx(MAIN_ORDER, True)
STATIC = L.Ctx(MAIN_ORDER, False)
assert len(MAIN_ORDER) <= 80

def setup_doc(landscape=False, double=False, size=12):
    doc = Document()
    st = doc.styles
    st["Normal"].font.name = FONT
    st["Normal"].font.size = Pt(size)
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st["Normal"].paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 2.0 if double else 1.0
    for name, ital in (("Heading 1", False), ("Heading 2", True)):
        s = st[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = ital
        s.font.color.rgb = RGBColor(0, 0, 0)
        rf = s.element.rPr.rFonts
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.space_after = Pt(0)
        s.paragraph_format.line_spacing = 2.0 if double else 1.0
        s.paragraph_format.keep_with_next = True
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    else:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    L.add_xml(fp, L.fld_runs("PAGE", "<w:r><w:t>1</w:t></w:r>"))
    return doc

def line_numbers(doc):
    sect = doc.sections[0]._sectPr
    ln = parse_xml('<w:lnNumType %s w:countBy="1" w:restart="continuous"/>' % nsdecls("w"))
    pgmar = sect.find(qn("w:pgMar"))
    pgmar.addnext(ln)

def para(doc, text="", ctx=None, align=None, indent=True, before_break=False, bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if before_break:
        p.paragraph_format.page_break_before = True
    if text:
        if ctx is None:
            L.add_text(p, text, size, bold=bold, italic=italic)
        else:
            ctx.fill(p, text, size)
    return p

def heading(doc, text, level=1, before_break=False):
    h = doc.add_heading(text, level=level)
    if before_break:
        h.paragraph_format.page_break_before = True
    return h

def write_core_props(doc, title):
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

def inject_zotero(tmp, final):
    pref = ('<data data-version="3" zotero-version="7.0.0"><session id="%s"/>'
            '<style id="%s" locale="en-US" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
            '<prefs><pref name="fieldType" value="Field"/></prefs></data>') % (L.rid(8), CSL_ID)
    custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ZOTERO_PREF_1"><vt:lpwstr>%s</vt:lpwstr></property></Properties>') % L.esc(pref)
    zin = zipfile.ZipFile(tmp)
    zout = zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            s = data.decode("utf-8")
            if "custom-properties" not in s:
                s = s.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
            data = s.encode("utf-8")
        elif item.filename == "_rels/.rels":
            s = data.decode("utf-8")
            if "custom-properties" not in s:
                s = s.replace("</Relationships>", '<Relationship Id="rIdCustom1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
            data = s.encode("utf-8")
        zout.writestr(item, data)
    zout.writestr("docProps/custom.xml", custom)
    zout.close()
    zin.close()
    os.remove(tmp)

def build_manuscript():
    doc = setup_doc(double=True)
    line_numbers(doc)
    ctx = MAIN
    p = para(doc, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    L.add_text(p, J.TITLE, 12, bold=True)
    p = para(doc, indent=False)
    for i, (name, sup) in enumerate(J.AUTHORS):
        L.add_text(p, name, 12)
        L.add_text(p, sup, 12).font.superscript = True
        if i < len(J.AUTHORS) - 1:
            L.add_text(p, " and ", 12)
    para(doc, J.DEGREES, ctx, indent=False)
    for num, txt in J.AFFIL:
        p = para(doc, indent=False)
        L.add_text(p, num, 12).font.superscript = True
        L.add_text(p, " " + txt, 12)
    para(doc, J.CORRESPONDENCE, ctx, indent=False)
    for label, txt in J.TITLE_PAGE_ITEMS:
        p = para(doc, indent=False)
        L.add_text(p, label + ". ", 12, italic=True)
        ctx.fill(p, txt, 12)
    heading(doc, "Abstract", 1, before_break=True)
    para(doc, J.ABSTRACT, None, indent=False)
    p = para(doc, indent=False)
    L.add_text(p, "Key words: ", 12, italic=True)
    L.add_text(p, J.KEYWORDS, 12)
    first = True
    heading(doc, "Introduction", 1, before_break=True)
    for t in J.INTRO:
        para(doc, t, ctx)
    heading(doc, "Methods", 1)
    for kind, t in J.METHODS:
        if kind == "h2":
            heading(doc, t, 2)
        else:
            para(doc, t, ctx)
    heading(doc, "Results", 1)
    for kind, t in J.RESULTS:
        if kind == "h2":
            heading(doc, t, 2)
        else:
            para(doc, t, ctx)
    heading(doc, "Discussion", 1)
    for t in J.DISCUSSION:
        para(doc, t, ctx)
    heading(doc, "Conclusions", 1)
    para(doc, J.CONCLUSIONS, ctx)
    heading(doc, "Acknowledgements", 1)
    para(doc, J.ACKNOWLEDGEMENTS, ctx)
    heading(doc, "References", 1, before_break=True)
    bibl = 'ADDIN ZOTERO_BIBL {&quot;uncited&quot;:[],&quot;omitted&quot;:[],&quot;custom&quot;:[]} CSL_BIBLIOGRAPHY'
    n = len(MAIN_ORDER)
    for i, k in enumerate(MAIN_ORDER):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        if i == 0:
            L.add_xml(p, '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> %s </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>' % bibl)
        L.add_text(p, "%d." % (i + 1), 12)
        p.add_run().add_tab()
        L.add_text(p, L.ref_text(k), 12)
        if i == n - 1:
            L.add_xml(p, '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
    heading(doc, "Figure Legends", 1, before_break=True)
    for label, txt in J.FIG_LEGENDS:
        p = para(doc, indent=False)
        L.add_text(p, label + ". ", 12, italic=True)
        L.add_text(p, txt, 12)
    heading(doc, "List of Supplemental Digital Content", 1, before_break=True)
    for num in sorted(J.SDC):
        desc, fname = J.SDC[num]
        para(doc, "Supplemental digital content %d" % num, indent=False)
        para(doc, desc, indent=False)
        para(doc, fname, indent=False)
    write_core_props(doc, J.TITLE)
    final = os.path.join(OUT, "Manuscript.docx")
    tmp = final + ".tmp"
    doc.save(tmp)
    inject_zotero(tmp, final)
    print("manuscript", final, "refs", n)

def cell_bottom(cell, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml('<w:tcBorders %s><w:bottom w:val="single" w:sz="%d" w:space="0" w:color="000000"/></w:tcBorders>' % (nsdecls("w"), sz)))

def table_borders(table):
    xml = ('<w:tblBorders %s><w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:left w:val="nil"/>'
           '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:right w:val="nil"/>'
           '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:insideV w:val="nil"/></w:tblBorders>') % nsdecls("w")
    table._tbl.tblPr.append(parse_xml(xml))

def cell_margins(table):
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", 40), ("left", 70), ("bottom", 40), ("right", 70)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    table._tbl.tblPr.append(m)

def add_table(doc, t, ctx, label, caption, note, fs, new_page=True, landscape=False):
    cp = doc.add_paragraph()
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_after = Pt(6)
    if new_page:
        cp.paragraph_format.page_break_before = True
    L.add_text(cp, label + ". ", 12, bold=True)
    L.add_text(cp, caption, 12)
    ncols = len(t["header"])
    table = doc.add_table(rows=1, cols=ncols)
    table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell_margins(table)
    sec = doc.sections[-1]
    uw = sec.page_width - sec.left_margin - sec.right_margin
    total = sum(t["widths"])
    widths = [int(uw * w / total) for w in t["widths"]]
    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for j, h in enumerate(t["header"]):
        c = hdr.cells[j]
        c.width = widths[j]
        cell_bottom(c, 8)
        pp = c.paragraphs[0]
        L.add_text(pp, h, fs, bold=True)
    for row in t["rows"]:
        r = table.add_row()
        r._tr.get_or_add_trPr().append(parse_xml('<w:cantSplit %s w:val="true"/>' % nsdecls("w")))
        if row[0] == "__GROUP__":
            merged = r.cells[0].merge(r.cells[ncols - 1])
            L.add_text(merged.paragraphs[0], row[1], fs + 0.5, bold=True, italic=True)
            merged.paragraphs[0].paragraph_format.keep_with_next = True
            continue
        for j, val in enumerate(row):
            c = r.cells[j]
            c.width = widths[j]
            ctx.fill(c.paragraphs[0], val, fs)
    np_ = doc.add_paragraph()
    np_.paragraph_format.space_before = Pt(6)
    L.add_text(np_, note, fs)
    return table

def build_tables():
    T_ = T.all_tables()
    t1 = dict(T_["t2"])
    t2 = dict(T_["t7"])
    rows2 = [list(r) for r in t2["rows"]]
    for r in rows2:
        if r[2] == "Table 5":
            r[2] = "Supplemental Digital Content 10"
    for r in rows2:
        r[1] = r[1].replace("Six supplementary studies", "Six instrument-focused studies")
        r[2] = r[2].replace("0 core, 6 supplementary", "0 core, 6 instrument-focused")
    t2["rows"] = rows2
    t1["widths"] = [1.25, 0.95, 3.0, 2.0, 2.0]
    note1 = ("Counts are studies, and a study can contribute to more than one row. CT, computed tomography. DWI, diffusion-weighted imaging. GCS, Glasgow Coma Scale. GOS, Glasgow Outcome Scale. "
             "GOSE, extended Glasgow Outcome Scale. KPS, Karnofsky Performance Status. mFI-5, five-factor modified frailty index. mFI-11, eleven-item modified frailty index. MEP, motor evoked potential. "
             "MMSE, Mini-Mental State Examination. MoCA, Montreal Cognitive Assessment. MoCA-INA, Indonesian version of the MoCA. mRS, modified Rankin Scale. NANO, Neurologic Assessment in Neuro-Oncology. "
             "NIHSS, National Institutes of Health Stroke Scale.")
    note2 = ("Status describes how fully the mapped evidence answers the question. Partly answered and gap findings are reported as results of a scoping review. DWI, diffusion-weighted imaging. "
             "GFAP, glial fibrillary acidic protein. MoCA-INA, Indonesian version of the Montreal Cognitive Assessment. NANO-Ina, Indonesian version of the Neurologic Assessment in Neuro-Oncology scale. "
             "NfL, neurofilament light chain. NSE, neuron-specific enolase. RQ, review question.")
    doc = setup_doc(size=12)
    first = True
    add_table(doc, t1, STATIC, "Table 1", "Candidate domains of tumor-associated brain frailty, instruments identified in the core evidence base, and remaining gaps.", note1, 10, new_page=False)
    add_table(doc, t2, STATIC, "Table 2", "Answers to the seven review questions and the status of the mapped evidence.", note2, 10)
    write_core_props(doc, "Tables")
    doc.save(os.path.join(OUT, "Tables.docx"))
    print("tables saved")

def sdc_doc(num, spec):
    doc = setup_doc(landscape=spec.get("landscape", False), size=11)
    rows_texts = []
    for t in spec["tables"]:
        for r in t["rows"]:
            rows_texts.extend(r)
    order = L.order_of(rows_texts + spec.get("extra_texts", []))
    ctx = L.Ctx(order, False)
    p = doc.add_paragraph()
    L.add_text(p, "Supplemental Digital Content %d" % num, 11, bold=True)
    p.paragraph_format.space_after = Pt(6)
    first = True
    for t, cap in zip(spec["tables"], spec["captions"]):
        add_table(doc, t, ctx, "Table", cap, t["note"], spec.get("fs", 9), new_page=not first)
        first = False
    if spec.get("figure"):
        fp = doc.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(12)
        fp.paragraph_format.keep_with_next = True
        fp.add_run().add_picture(os.path.join(FIGS, spec["figure"][0]), width=Inches(spec["figure"][1]))
        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(4)
        L.add_text(lp, "Figure. ", 11, bold=True)
        L.add_text(lp, spec["figure"][2], 11)
    if order:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.keep_with_next = True
        L.add_text(h, "References", 11, bold=True)
        for i, k in enumerate(order):
            rp = doc.add_paragraph()
            rp.paragraph_format.left_indent = Inches(0.4)
            rp.paragraph_format.first_line_indent = Inches(-0.4)
            rp.paragraph_format.space_after = Pt(2)
            L.add_text(rp, "%d.\t%s" % (i + 1, ref_text_sdc(k)), 9)
    write_core_props(doc, "Supplemental Digital Content %d" % num)
    doc.save(os.path.join(OUT, "SDC %d.docx" % num))
    print("SDC", num)

def ref_text_sdc(k):
    if k == "core60":
        r = L.REFS[k]
        au = ["%s %s" % (a["family"], a["ini"]) for a in r["authors"]]
        authors = ", ".join(au[:3]) + (", et al." if len(au) > 3 else ".")
        return "%s %s arXiv. Published online 2025. DOI: %s" % (authors, L.lower_after_colon(r["title"]).rstrip(".") + ".", r["doi"])
    return L.ref_text(k)

def build_sdc_docs():
    T_ = T.all_tables()
    specs = {
        6: {"tables": [T_["t1"]], "captions": ["Characteristics of the 70 studies in the core evidence base, grouped by evidence stream and ordered by year within stream."],
            "landscape": True, "fs": 7, "figure": ("jna_sdc_characteristics.png", 6.2, "Characteristics of the 70 core studies. (A) Publication period. (B) Study design. (C) Tumor category of the study sample. (D) Evidence stream.")},
        7: {"tables": [], "captions": [], "figure": ("jna_sdc_framework.png", 6.3, "Four candidate domains of tumor-associated brain frailty, with the number of core studies, instruments identified, and gaps, above the shared candidate mechanism of neuroinflammation and barrier injury. IL, interleukin. TNF, tumor necrosis factor. VEGF, vascular endothelial growth factor. NSE, neuron-specific enolase. GFAP, glial fibrillary acidic protein. NfL, neurofilament light chain. mFI, modified frailty index. MMSE, Mini-Mental State Examination. MoCA, Montreal Cognitive Assessment. KPS, Karnofsky Performance Status. NIHSS, National Institutes of Health Stroke Scale. mRS, modified Rankin Scale. GOS, Glasgow Outcome Scale. NANO, Neurologic Assessment in Neuro-Oncology. DWI, diffusion-weighted imaging.")},
        8: {"tables": [T_["t3"]], "captions": ["Neuroinflammation, injury, and vulnerability evidence by tumor category, from the 70 core studies (protocols excluded)."], "fs": 8},
        9: {"tables": [T_["t4"]], "captions": ["Biomarkers measured in the core evidence base, with source, interpretation, sampling, and main findings."], "fs": 8},
        10: {"tables": [T_["t5"]], "captions": ["Evidence on perioperative brain injury as a second hit after tumor-related vulnerability."], "fs": 9},
        11: {"tables": [T_["t6"]], "captions": ["Instrument-focused stratum, studies of MoCA-INA, NANO-Ina, and their parent instruments."], "fs": 8},
    }
    for num, spec in specs.items():
        sdc_doc(num, spec)

CHECKLIST_WHERE = {
    2: "Abstract (unstructured, as the journal requires)",
    3: "Introduction",
    4: "Introduction (RQ1 to RQ7) and Methods, Eligibility Criteria",
    5: "Methods, Review Design and Reporting (not registered, materials deposited on GitHub and archived on Zenodo)",
    6: "Methods, Eligibility Criteria",
    7: "Methods, Search Strategy (last searches March 22, 2026)",
    8: "Methods, Search Strategy and Supplemental Digital Content 2",
    9: "Methods, Study Selection and Data Charting (two reviewers screened independently, disagreements resolved in a video meeting)",
    10: "Methods, Study Selection and Data Charting and Supplemental Digital Content 4",
    11: "Methods, Study Selection and Data Charting",
    12: "Methods, Evidence Classification and Synthesis (not performed, design-based evidence level used)",
    13: "Methods, Evidence Classification and Synthesis",
    14: "Results, Study Selection and Figure 1",
    15: "Results, Characteristics of the Core Evidence Base and Supplemental Digital Content 6",
    17: "Supplemental Digital Content 4 and Supplemental Digital Content 6",
    18: "Results (all subsections), Tables 1 and 2, Figures 2 and 3, Supplemental Digital Content 7 to 11",
    19: "Discussion",
    20: "Discussion",
    21: "Conclusions",
    22: "Title page (no external funding)",
}
XLSX_SHEETS = {1: "S1 PRISMA-ScR", 2: "S2 Searches", 3: "S4 Reconciliation", 4: "S3 Charting form", 5: "S6 Coding"}

def build_sdc_xlsx():
    src = "../outputs/Supplementary_Tables_S1-S6.xlsx"
    for num, sheet in XLSX_SHEETS.items():
        wb = openpyxl.load_workbook(src)
        for name in list(wb.sheetnames):
            if name != sheet:
                del wb[name]
        ws = wb[sheet]
        ws.title = "SDC %d" % num
        if num == 1:
            for row in ws.iter_rows(min_row=2):
                item = row[0].value
                if item in CHECKLIST_WHERE:
                    row[3].value = CHECKLIST_WHERE[item]
        wb.save(os.path.join(OUT, "SDC %d.xlsx" % num))
        print("SDC xlsx", num)

if __name__ == "__main__":
    build_manuscript()
    build_tables()
    build_sdc_docs()
    build_sdc_xlsx()
