from docx.enum.text import WD_COLOR_INDEX
import re, json, sys, random, string, zipfile, shutil, os
from xml.sax.saxutils import escape
sys.stdout.reconfigure(encoding="utf-8")
BUILD = "."
sys.path.insert(0, BUILD)
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import assemble as AS
import content_c as C
import tables as T

OUT = "../outputs/Scoping_Review_PBI_Tumor_Associated_Brain_Frailty.docx"
UID = 14681764
random.seed(20261017)
REFS = json.load(open(BUILD + r"\refs_master.json", encoding="utf-8"))
ZK = json.load(open(BUILD + "/zkeys.json"))["keys"]
ORDER = json.load(open(BUILD + r"\cite_order.json"))
CNUM = {k: i + 1 for i, k in enumerate(ORDER)}
TOKEN = AS.TOKEN
FONT = "Times New Roman"

def rid(n=8):
    return "".join(random.choice(string.ascii_lowercase + string.digits) for _ in range(n))

def esc(s):
    return escape(s, {'"': "&quot;"})

def csl_item(k):
    r = REFS[k]
    n = CNUM[k]
    d = {"id": n}
    if k == "core60":
        d.update({"type": "article", "genre": "preprint", "title": r["title"], "container-title": "arXiv", "DOI": r["doi"], "number": "arXiv:2507.23057", "URL": "https://arxiv.org/abs/2507.23057"})
    elif k in ("bangun2018", "mckenzie2019"):
        d.update({"type": "chapter", "title": r["title"], "container-title": "Stem Cell Oncology" if k == "bangun2018" else "Cochrane Handbook for Systematic Reviews of Interventions",
                  "publisher": "CRC Press" if k == "bangun2018" else "Wiley", "page": r["pages"], "DOI": r["doi"]})
    else:
        d.update({"type": "article-journal", "title": r["title"], "container-title": r["journal"], "container-title-short": r["abbr"], "journalAbbreviation": r["abbr"]})
        if r["volume"]:
            d["volume"] = r["volume"]
        if r["issue"]:
            d["issue"] = r["issue"]
        if r["pages"]:
            d["page"] = r["pages"]
        if r["doi"]:
            d["DOI"] = r["doi"]
        if r.get("pmid"):
            d["PMID"] = r["pmid"]
    d["issued"] = {"date-parts": [[int(r["year"])]]}
    d["language"] = "en"
    d["author"] = [{"family": a["family"], "given": a["given"]} for a in r["authors"]]
    return {"id": n, "uris": ["http://zotero.org/users/%d/items/%s" % (UID, ZK[k])], "itemData": d}

def collapse(nums):
    nums = sorted(set(nums))
    out = []
    i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            out.append("%d-%d" % (nums[i], nums[j]))
        else:
            out.extend(str(x) for x in nums[i:j + 1])
        i = j + 1
    return ",".join(out)

def fld_runs(instr, result_xml):
    return ('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> %s </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>%s'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r>') % (instr, result_xml)

def cite_runs(keylist, rpr_extra=""):
    keys = [AS.rk(k) for k in keylist.split(",")]
    nums = [CNUM[k] for k in keys]
    shown = collapse(nums)
    props = {"unsorted": False, "formattedCitation": "\\super %s\\nosupersub{}" % shown, "plainCitation": shown, "noteIndex": 0}
    cj = {"citationID": rid(8), "properties": props, "citationItems": [csl_item(k) for k in sorted(keys, key=lambda x: CNUM[x])],
          "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"}
    instr = "ADDIN ZOTERO_ITEM CSL_CITATION " + esc(json.dumps(cj, ensure_ascii=False, separators=(",", ":")))
    res = '<w:r><w:rPr><w:vertAlign w:val="superscript"/>%s</w:rPr><w:t>%s</w:t></w:r>' % (rpr_extra, esc(shown))
    return fld_runs(instr, res)

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

def add_xml(par, xml):
    for el in parse_xml("<w:p %s>%s</w:p>" % (W_NS, xml)):
        par._p.append(el)

def add_text(par, text, size=None, bold=False, italic=False):
    r = par.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return r

HL = re.compile(r"\{hl:([^}]+)\}")

def fill(par, text, size=None):
    pieces = HL.split(text)
    for idx, piece in enumerate(pieces):
        if idx % 2 == 0:
            if piece:
                fill_plain(par, piece, size)
        else:
            r = add_text(par, piece, size)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def fill_plain(par, text, size=None):
    pos = 0
    for m in TOKEN.finditer(text):
        seg = text[pos:m.start()].rstrip(" ")
        if seg:
            add_text(par, seg, size)
        rpr = '<w:sz w:val="%d"/>' % int(size * 2) if size else ""
        add_xml(par, cite_runs(m.group(1), rpr))
        pos = m.end()
    tail = text[pos:]
    if tail:
        add_text(par, tail, size)

def minimal_range(p):
    m = re.fullmatch(r"(\d+)-(\d+)", p)
    if not m:
        return p
    a, b = m.group(1), m.group(2)
    if len(a) == len(b):
        i = 0
        while i < len(a) - 1 and a[i] == b[i]:
            i += 1
        return a + "-" + b[i:]
    return p

def ref_text(k):
    r = REFS[k]
    au = ["%s %s" % (a["family"], a["ini"]) for a in r["authors"]]
    if len(au) > 6:
        au = au[:6] + ["et al"]
    authors = ", ".join(au)
    from zupload import lowercase_after_colon
    title = lowercase_after_colon(r["title"]).rstrip(".")
    t = title if title[-1] in "?!" else title + "."
    if k == "core60":
        return "%s. %s arXiv. 2025. doi:%s" % (authors, t, r["doi"])
    if k == "bangun2018":
        return "%s. %s In: Stem Cell Oncology. London: CRC Press; 2018. p. 119-23. doi:%s" % (authors, t, r["doi"])
    if k == "mckenzie2019":
        return "%s. %s In: Higgins JPT, Thomas J, Chandler J, Cumpston M, Li T, Page MJ, et al., editors. Cochrane Handbook for Systematic Reviews of Interventions. 2nd ed. Chichester: Wiley; 2019. p. 321-47. doi:%s" % (authors, t, r["doi"])
    s = "%s. %s %s. %s" % (authors, t, r["abbr"], r["year"])
    loc = ""
    if r["volume"]:
        loc += r["volume"]
    if r["issue"]:
        loc += "(%s)" % r["issue"]
    if r["pages"]:
        loc += ":" + minimal_range(r["pages"])
    if loc:
        s += ";" + loc
    s += "."
    if r["doi"]:
        s += " doi:" + r["doi"]
    return s

doc = Document()
st = doc.styles
st["Normal"].font.name = FONT
st["Normal"].font.size = Pt(11)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st["Normal"].paragraph_format
pf.space_after = Pt(6)
pf.line_spacing = 1.15
for name, sz, bold, ital, before, after in (("Heading 1", 13, True, False, 14, 6), ("Heading 2", 11, True, True, 10, 4)):
    s = st[name]
    s.font.name = FONT
    s.font.size = Pt(sz)
    s.font.bold = bold
    s.font.italic = ital
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    s.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if s.element.rPr.rFonts.get(qn(a)) is not None:
            del s.element.rPr.rFonts.attrib[qn(a)]
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
cap = st["Caption"]
cap.font.name = FONT
cap.font.size = Pt(9.5)
cap.font.bold = False
cap.font.italic = False
cap.font.color.rgb = RGBColor(0, 0, 0)
cap.paragraph_format.space_after = Pt(4)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.3))

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_xml(fp, fld_runs("PAGE", '<w:r><w:t>1</w:t></w:r>'))
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(hp, "Tumor-associated brain frailty, a scoping review", 8.5, italic=True)

TABLES = T.all_tables()
TNUM = {"t%d" % i: i for i in range(1, 8)}
FNUM = {"fig%d" % i: i for i in range(1, 6)}
D = AS.doc_blocks()

def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)

def set_borders(table):
    tblPr = table._tbl.tblPr
    xml = ('<w:tblBorders %s><w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:left w:val="nil"/>'
           '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/><w:right w:val="nil"/>'
           '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/><w:insideV w:val="nil"/></w:tblBorders>') % nsdecls("w")
    tblPr.append(parse_xml(xml))

def cell_bottom(cell, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml('<w:tcBorders %s><w:bottom w:val="single" w:sz="%d" w:space="0" w:color="000000"/></w:tcBorders>' % (nsdecls("w"), sz)))

def usable_width(section):
    return section.page_width - section.left_margin - section.right_margin

def add_table_block(key):
    t = TABLES[key]
    n = TNUM[key]
    landscape = t["landscape"]
    if landscape:
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin"):
            setattr(s, side, Cm(1.8))
        s.top_margin = s.bottom_margin = Cm(2.0)
    cp = doc.add_paragraph(style="Caption")
    cp.paragraph_format.keep_with_next = True
    label, rest = t["caption"].split(". ", 1)
    add_text(cp, label + ". ", 9.5, bold=True)
    add_text(cp, rest, 9.5)
    ncols = len(t["header"])
    table = doc.add_table(rows=1, cols=ncols)
    set_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_margins(table)
    total = sum(t["widths"])
    uw = usable_width(doc.sections[-1])
    widths = [int(uw * w / total) for w in t["widths"]]
    fs = t["fs"]
    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for j, h in enumerate(t["header"]):
        c = hdr.cells[j]
        c.width = widths[j]
        cell_bottom(c, 8)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_text(p, h, fs, bold=True)
    for row in t["rows"]:
        r = table.add_row()
        trPr = r._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)
        if row[0] == "__GROUP__":
            merged = r.cells[0].merge(r.cells[ncols - 1])
            p = merged.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            add_text(p, row[1], fs + 0.5, bold=True, italic=True)
            continue
        for j, val in enumerate(row):
            c = r.cells[j]
            c.width = widths[j]
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            fill(p, val, fs)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    add_text(note, t["note"], 8.5)
    if landscape:
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, side, Cm(2.3))

def add_figure_block(key):
    pic = os.path.join(BUILD, "figs", C.FIGFILE[key])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(pic, width=Cm(C.FIGWIDTH[key]))
    cp = doc.add_paragraph(style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label, rest = C.CAPTIONS[key].split(". ", 1)
    add_text(cp, label + ". ", 9.5, bold=True)
    add_text(cp, rest, 9.5)

for kind, payload in D:
    if kind == "title":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        add_text(p, payload, 16, bold=True)
    elif kind == "authors":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        for txt, sup in payload["names"]:
            r = add_text(p, txt, 11.5, bold=True)
            r.font.superscript = bool(sup)
        for a in payload["affil"]:
            q = doc.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(0)
            q.paragraph_format.line_spacing = 1.0
            r = add_text(q, a[0], 9)
            r.font.superscript = True
            add_text(q, " " + a[1], 9)
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = Pt(4)
        add_text(q, payload["corr"], 9)
    elif kind == "abstract":
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        add_text(h, "Abstract", 12, bold=True)
        for label, text in payload:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_text(p, label + " ", 10.5, bold=True)
            add_text(p, text, 10.5)
    elif kind == "keywords":
        p = doc.add_paragraph()
        add_text(p, "Keywords. ", 10.5, bold=True)
        add_text(p, payload, 10.5)
    elif kind == "h1":
        doc.add_heading(payload, level=1)
    elif kind == "h1b":
        doc.add_heading(payload, level=2)
    elif kind == "h2":
        doc.add_heading(payload, level=2)
    elif kind == "p":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fill(p, payload)
    elif kind == "table":
        add_table_block(payload)
    elif kind == "figure":
        add_figure_block(payload)

doc.add_heading("References", level=1)
bibl_instr = 'ADDIN ZOTERO_BIBL {&quot;uncited&quot;:[],&quot;omitted&quot;:[],&quot;custom&quot;:[]} CSL_BIBLIOGRAPHY'
n = len(ORDER)
for i, k in enumerate(ORDER):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    if i == 0:
        add_xml(p, '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> %s </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>' % bibl_instr)
    add_text(p, "%d." % (i + 1), 9.5)
    p.add_run().add_tab()
    add_text(p, ref_text(k), 9.5)
    if i == n - 1:
        add_xml(p, '<w:r><w:fldChar w:fldCharType="end"/></w:r>')

doc.core_properties.title = "Tumor-Associated Brain Frailty and Perioperative Brain Injury in Neuro-Oncologic Surgery: A Scoping Review"
doc.core_properties.author = ""
doc.core_properties.comments = "Citations are live Zotero fields (NLM superscript). Use Zotero > Refresh in Word to update."
tmp = OUT + ".tmp"
doc.save(tmp)

pref = ('<data data-version="3" zotero-version="7.0.0"><session id="%s"/>'
        '<style id="http://www.zotero.org/styles/nlm-citation-sequence-superscript" locale="en-US" hasBibliography="1" bibliographyStyleHasBeenSet="1"/>'
        '<prefs><pref name="fieldType" value="Field"/></prefs></data>') % rid(8)
custom = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
          '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
          '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ZOTERO_PREF_1"><vt:lpwstr>%s</vt:lpwstr></property></Properties>') % esc(pref)
zin = zipfile.ZipFile(tmp)
zout = zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == "[Content_Types].xml":
        s = data.decode("utf-8")
        if "custom-properties" not in s:
            s = s.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
        data = s.encode("utf-8")
    elif item.filename == "_rels/.rels":
        s = data.decode("utf-8")
        if "custom-properties" not in s:
            s = s.replace("</Relationships>", '<Relationship Id="rIdCustom1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        data = s.encode("utf-8")
    zout.writestr(item, data)
zout.writestr("docProps/custom.xml", custom)
zout.close()
zin.close()
os.remove(tmp)
print("saved", OUT, "refs", n)
